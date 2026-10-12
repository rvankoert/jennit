import asyncio
from datetime import datetime

import openpyxl

from jennit import report


def _fake_row():
    return {
        "rel_path": "a/b.jpg", "folder_path": "a", "status": "done",
        "finished_at": datetime(2026, 9, 29, 12, 0), "is_valid": True,
        "is_overexposed": False, "file_size_bytes": 123, "md5": "m",
        "sha512": "s", "dpi": [300, 300], "compression_ratio": 12.5,
        "jpeg_quality": 92, "black_pixels": 10, "white_pixels": 5,
        "angle": 0.0, "color_profile": None, "color_histograms": None,
        "exif": "{}", "errors": None,
    }


def _data():
    async def get_folder_rows(name):
        return {"(root)": [_fake_row()], "=CMD()": []}[name]

    return {
        "totals": {"total": 2, "valid": 2, "avg_jpeg_quality": 92.0},
        "folders": [
            {"folder": "(root)", "total": 1, "valid": 1, "avg_angle": 0.0,
             "avg_compression": 12.5, "avg_jpeg_quality": 92.0,
             "min_jpeg_quality": 92, "max_jpeg_quality": 92,
             "black_sum": 10, "white_sum": 5, "overexposed_count": 0,
             "avg_file_size": 123.0, "error_count": 0},
            {"folder": "=CMD()", "total": 1, "valid": 1, "avg_angle": None,
             "avg_compression": None, "avg_jpeg_quality": None,
             "min_jpeg_quality": None, "max_jpeg_quality": None,
             "black_sum": 0, "white_sum": 0, "overexposed_count": 0,
             "avg_file_size": None, "error_count": 0},
        ],
        "get_folder_rows": get_folder_rows,
    }


def _batch():
    return {"name": "t", "status": "completed",
            "created_at": datetime(2026, 9, 29, 11, 0),
            "completed_at": datetime(2026, 9, 29, 12, 30)}


def test_workbook_structure_and_formula_injection(tmp_path):
    out = tmp_path / "r.xlsx"
    asyncio.run(report.build_workbook(_data(), _batch(), out))

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Summary", "(root)", "=CMD()"], wb.sheetnames
    s = wb["Summary"]
    assert s["A6"].value == "Scans" and s["B6"].value == 2
    # I-006: a folder named '=CMD()' must be stored as literal text
    rows = [r for r in s.iter_rows(values_only=True) if r[0] == "'=CMD()"]
    assert rows
    d = wb["(root)"]
    assert [c.value for c in d[4]][0] == "Rel Path"


def test_sheet_name_limits():
    assert report._sheet_name("") == "(root)"
    assert report._sheet_name("a" * 100) == "a" * 31
    assert report._sheet_name("a/b\\c:d*e?f[g]h") == "abcdefgh"
