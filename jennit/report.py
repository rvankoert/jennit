"""Batch reports.

The workbook is a summary (one row per folder) plus one detail sheet per
folder, because Excel's hard row limit (1,048,576) makes per-scan rows
unsafe at the 1,000,000-scan scale. Full per-scan data is exported via
export_batch_csv / export_batch_parquet, streamed from Postgres.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from . import db
from .analysis import formula_safe
from .config import Settings

_BOLD = Font(bold=True)


def _fmt(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value)


async def build_workbook(pool_rows: dict, batch: dict, out_path: Path) -> Path:
    """Build the batch workbook from pre-fetched query results.

    pool_rows keys: totals (total/valid/avg_jpeg_quality), folders
    (per-folder aggregate rows), get_folder_rows (async callable:
    folder name -> per-scan rows; fetched one folder at a time so
    memory stays bounded). Kept DB-free so it is unit-testable without
    a database.
    """
    wb = Workbook()
    totals = pool_rows["totals"]
    total = totals["total"] or 0
    valid = totals["valid"] or 0
    invalid = total - valid
    quality_avg = totals["avg_jpeg_quality"]

    # --- sheet 1: summary (one row per folder) ---
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Jennit Batch Report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Batch", batch["name"]])
    ws.append(["Status", batch["status"]])
    ws.append(["Created", str(batch["created_at"])])
    ws.append(["Completed", str(batch["completed_at"]) if batch["completed_at"] else ""])
    ws.append(["Scans", total])
    ws.append(["Valid scans", valid])
    ws.append(["Invalid scans", invalid])
    ws.append(["Avg JPEG quality", _fmt(quality_avg)])
    ws.append([])
    ws.append(["Folder", "Scans", "Valid", "Invalid", "Avg angle (deg)",
               "Avg compression ratio", "Min JPEG quality", "Max JPEG quality",
               "Black pixels (sum)", "White pixels (sum)", "Overexposed",
               "Avg file size (bytes)", "Errors"])
    for cell in ws[ws.max_row]:
        cell.font = _BOLD
    folder_names: list[str] = []
    for f in pool_rows["folders"]:
        folder_names.append(f["folder"])
        ws.append([
            formula_safe(f["folder"]),
            f["total"],
            f["valid"] or 0,
            (f["total"] or 0) - (f["valid"] or 0),
            _fmt(f["avg_angle"]),
            _fmt(f["avg_compression"]),
            _fmt(f["min_jpeg_quality"]),
            _fmt(f["max_jpeg_quality"]),
            f["black_sum"] or 0,
            f["white_sum"] or 0,
            f["overexposed_count"] or 0,
            _fmt(f["avg_file_size"]),
            f["error_count"] or 0,
        ])
    # freeze the summary header row and widen columns a little
    # (in this fixed layout the summary header is always row 11)
    ws.freeze_panes = "A12"
    for col, width in (("A", 40), ("B", 10), ("M", 60)):
        ws.column_dimensions[col].width = width
    # --- sheets 2..n: one per folder (per-scan rows, capped) ---
    get_folder_rows = pool_rows["get_folder_rows"]
    for name in folder_names:
        rows = await get_folder_rows(name)
        ws_f = wb.create_sheet(_sheet_name(name))
        ws_f.append([f"Folder: {name}"])
        ws_f["A1"].font = _BOLD
        note = f"Per-scan rows: {len(rows)}"
        if len(rows) > FOLDER_DETAIL_LIMIT:
            note += (f" (sheet capped at {FOLDER_DETAIL_LIMIT}; "
                     "use the CSV export for all rows)")
        ws_f.append([note])
        ws_f.append([])
        ws_f.append(list(SCAN_HEADERS))
        for cell in ws_f[ws_f.max_row]:
            cell.font = _BOLD
        for r in rows[:FOLDER_DETAIL_LIMIT]:
            rec = _row_to_record(r)
            ws_f.append([_record_cell(rec[h]) for h in SCAN_HEADERS])
        ws_f.freeze_panes = "A5"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


FOLDER_DETAIL_LIMIT = 10_000


def _sheet_name(folder: str) -> str:
    r"""Excel sheet name: at most 31 chars, no []:*?/\ characters."""
    name = "".join(c for c in (folder or "") if c not in "[]:*?/\\")
    return name[:31].strip() or "(root)"


SCAN_HEADERS = ["Rel Path", "Status", "Finished", "Valid", "Overexposed",
                "File Size (bytes)", "MD5", "SHA-512", "DPI",
                "Compression Ratio", "JPEG Quality", "Black Pixels",
                "White Pixels", "Angle (deg)", "Color Profile",
                "Color Histograms", "EXIF", "Errors"]


def _row_to_record(row: dict) -> dict:
    """Map one scans-table row (db dict) to a flat record keyed by SCAN_HEADERS."""
    dpi = row.get("dpi")

    def num(v):
        return None if v is None else float(v)

    def i64(v):
        return None if v is None else int(v)

    return {
        "Rel Path": row["rel_path"],
        "Status": row["status"],
        "Finished": str(row["finished_at"]) if row.get("finished_at") else None,
        "Valid": None if row["is_valid"] is None else str(row["is_valid"]),
        "Overexposed": None if row["is_overexposed"] is None
            else str(row["is_overexposed"]),
        "File Size (bytes)": i64(row["file_size_bytes"]),
        "MD5": row["md5"],
        "SHA-512": row["sha512"],
        "DPI": f"{dpi[0]}x{dpi[1]}" if dpi else None,
        "Compression Ratio": num(row["compression_ratio"]),
        "JPEG Quality": i64(row["jpeg_quality"]),
        "Black Pixels": i64(row["black_pixels"]),
        "White Pixels": i64(row["white_pixels"]),
        "Angle (deg)": num(row["angle"]),
        "Color Profile": row["color_profile"],
        "Color Histograms": row["color_histograms"],
        "EXIF": row["exif"],
        "Errors": row["errors"],
    }


def _record_cell(value):
    """Cell value for a sheet: None -> '', strings formula-safe, floats rounded."""
    if value is None:
        return ""
    if isinstance(value, str):
        return formula_safe(value)
    if isinstance(value, float):
        return round(value, 2)
    return value


async def export_batch_csv(pool, batch_id, out_path: Path,
                           folder: str | None = None) -> int:
    """Stream every scan of the batch to CSV. Returns the number of data rows."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with open(out_path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(SCAN_HEADERS)
        async for row in db.iter_export_rows(pool, batch_id, folder=folder):
            rec = _row_to_record(row)
            writer.writerow(["" if rec[h] is None else rec[h] for h in SCAN_HEADERS])
            count += 1
    return count


async def export_batch_parquet(pool, batch_id, out_path: Path,
                               folder: str | None = None) -> int:
    try:
        import pyarrow as pa
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise RuntimeError(
            "parquet export requires pyarrow: pip install pyarrow") from exc

    types = {
        "Rel Path": pa.string(), "Status": pa.string(),
        "Finished": pa.string(), "Valid": pa.string(),
        "Overexposed": pa.string(), "File Size (bytes)": pa.int64(),
        "MD5": pa.string(), "SHA-512": pa.string(), "DPI": pa.string(),
        "Compression Ratio": pa.float64(), "JPEG Quality": pa.int64(),
        "Black Pixels": pa.int64(), "White Pixels": pa.int64(),
        "Angle (deg)": pa.float64(), "Color Profile": pa.string(),
        "Color Histograms": pa.string(), "EXIF": pa.string(),
        "Errors": pa.string(),
    }
    schema = pa.schema([(h, types[h]) for h in SCAN_HEADERS])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    writer = pq.ParquetWriter(out_path, schema)
    buf: list[dict] = []

    def flush() -> None:
        nonlocal count
        if not buf:
            return
        cols = {h: [rec[h] for rec in buf] for h in SCAN_HEADERS}
        writer.write_table(pa.table(cols, schema=schema))
        count += len(buf)
        buf.clear()

    async for row in db.iter_export_rows(pool, batch_id, folder=folder):
        buf.append(_row_to_record(row))
        if len(buf) >= 10_000:
            flush()
    flush()
    writer.close()
    return count


async def write_report(pool, settings: Settings, batch_id) -> Path:
    """Generate (or regenerate) the batch workbook and store its path on the batch."""
    batch = await db.get_batch(pool, batch_id)
    if batch is None:
        raise ValueError(f"unknown batch: {batch_id}")
    totals = await pool.fetchrow(
        """
        SELECT count(*) AS total,
               count(*) FILTER (WHERE is_valid = TRUE) AS valid,
               avg(jpeg_quality) AS avg_jpeg_quality
        FROM scans WHERE batch_id = $1
        """,
        batch_id,
    )
    folders = [dict(r) for r in await pool.fetch(
        """
        SELECT COALESCE(NULLIF(folder_path, ''), '(root)') AS folder,
               count(*) AS total,
               count(*) FILTER (WHERE is_valid = TRUE) AS valid,
               avg(angle) FILTER (WHERE angle IS NOT NULL) AS avg_angle,
               avg(compression_ratio)
                   FILTER (WHERE compression_ratio IS NOT NULL) AS avg_compression,
               avg(jpeg_quality) AS avg_jpeg_quality,
               min(jpeg_quality) AS min_jpeg_quality,
               max(jpeg_quality) AS max_jpeg_quality,
               coalesce(sum(black_pixels), 0) AS black_sum,
               coalesce(sum(white_pixels), 0) AS white_sum,
               count(*) FILTER (WHERE is_overexposed = TRUE) AS overexposed_count,
               avg(file_size_bytes) AS avg_file_size,
               count(*) FILTER (WHERE errors IS NOT NULL) AS error_count
        FROM scans WHERE batch_id = $1
        GROUP BY folder ORDER BY folder
        """,
        batch_id,
    )]

    async def get_folder_rows(name: str) -> list[dict]:
        rows = await pool.fetch(
            "SELECT * FROM scans WHERE batch_id = $1 AND folder_path = $2 "
            "ORDER BY id LIMIT $3",
            batch_id, "" if name == "(root)" else name, FOLDER_DETAIL_LIMIT + 1,
        )
        return [dict(r) for r in rows]

    out_path = settings.data_root / "reports" / str(batch_id) / "batch_report.xlsx"
    await build_workbook(
        {"totals": dict(totals), "folders": folders,
         "get_folder_rows": get_folder_rows},
        batch, out_path,
    )
    await pool.execute(
        "UPDATE batches SET report_path = $2 WHERE id = $1",
        batch_id, str(out_path),
    )
    return out_path
